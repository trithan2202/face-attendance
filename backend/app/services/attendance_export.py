import calendar
import io
from datetime import date, datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from attendance.attendance_calculator import AttendanceCalculator
from backend.app.core.config import settings
from backend.app.repositories.attendance_repository import AttendanceRepository
from backend.app.services.settings_service import SettingsService


XLSX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)

THIN_BLACK = Side(style="thin", color="000000")
GRID_BORDER = Border(
    left=THIN_BLACK,
    right=THIN_BLACK,
    top=THIN_BLACK,
    bottom=THIN_BLACK,
)
HEADER_FILL = PatternFill("solid", fgColor="A6A6A6")
SUBHEADER_FILL = PatternFill("solid", fgColor="B7B7B7")
WORKDAY_FILL = PatternFill("solid", fgColor="FFF200")
PRESENT_FILL = PatternFill("solid", fgColor="57D16C")
WARNING_FILL = PatternFill("solid", fgColor="F8CBAD")
ABSENT_FILL = PatternFill("solid", fgColor="F4CCCC")
LEAVE_FILL = PatternFill("solid", fgColor="BDD7EE")
WEEKEND_FILL = PatternFill("solid", fgColor="FF4D4D")
INVALID_DAY_FILL = PatternFill("solid", fgColor="D9D9D9")
OT_INFO_FILL = PatternFill("solid", fgColor="5B9BD5")
OT_DATE_GROUP_FILL = PatternFill("solid", fgColor="FFC000")
OT_DATE_FILL = PatternFill("solid", fgColor="FF0000")
OT_TOTAL_FILL = PatternFill("solid", fgColor="82DC8C")
OT_NOTE_FILL = PatternFill("solid", fgColor="F2F2F2")
# Dòng Giờ vào/Giờ ra tô xám để dòng Giờ OT (dòng được cộng) nổi bật.
OT_PUNCH_FILL = PatternFill("solid", fgColor="BFBFBF")
NOTE_PAID_FILL = PatternFill("solid", fgColor="8DB4E2")
NOTE_PAID_ROW_FILL = PatternFill("solid", fgColor="C5D9F1")
NOTE_INSURANCE_FILL = PatternFill("solid", fgColor="A6A6A6")
NOTE_INSURANCE_ROW_FILL = PatternFill("solid", fgColor="D9D9D9")

LOGO_PATH = Path(__file__).resolve().parents[1] / "assets" / "logo_ACACY.png"
# Logo cao 79px, vừa khít 4 dòng cao 15pt (20px) ở góc trên bên trái.
LOGO_ROWS = 4


def _safe_text(value):
    text = "" if value is None else str(value)
    return f"'{text}" if text.startswith(("=", "+", "-", "@")) else text


def _department_name(employee):
    return _safe_text(
        employee.department.name if employee.department else None
    )


def _shift_minutes(shift):
    if shift is None:
        return 0
    start = datetime.combine(date.min, shift.start_time)
    end = datetime.combine(date.min, shift.end_time)
    if shift.cross_day or end <= start:
        end += timedelta(days=1)
    return max(0, int((end - start).total_seconds() // 60))


def _workday_code(daily, shift, half_day_ratio=0.75):
    if daily is None:
        return "", None
    if daily.status == "ABSENT":
        return "V", 0
    if daily.status == "MISSING_CHECK_IN":
        return "TCI", 0
    if daily.status == "MISSING_CHECK_OUT":
        return "TCO", 0
    if daily.status == "ON_LEAVE":
        return "NP", 0
    if daily.status != "PRESENT":
        return _safe_text(daily.status), 0

    scheduled_minutes = _shift_minutes(shift)
    if scheduled_minutes and daily.worked_minutes < scheduled_minutes * half_day_ratio:
        return "0.5HC", 0.5
    return "HC", 1


def _status_label(status):
    return {
        "PRESENT": "Có mặt",
        "MISSING_CHECK_IN": "Thiếu check-in",
        "MISSING_CHECK_OUT": "Thiếu check-out",
        "ABSENT": "Vắng",
        "ON_LEAVE": "Nghỉ phép",
    }.get(status, status)


def _overtime_hours(minutes):
    """Quy giờ tăng ca về số giờ tròn.

    Phần lẻ dưới 30 phút thì làm tròn xuống, từ 30 phút trở lên làm tròn lên:
    1:15 thành 1 giờ, 1:45 thành 2 giờ.
    """
    return (max(0, int(minutes or 0)) + 30) // 60


def _report_date():
    # Container chạy giờ UTC; ngày xuất phải theo giờ công ty.
    return datetime.now(ZoneInfo(settings.TIMEZONE)).date()


def _clock(value):
    """Giờ trong ngày (kiểu time) để Excel hiển thị theo định dạng h:mm."""
    return value.time().replace(second=0, microsecond=0) if value else None


def _overtime_status_label(status):
    return {
        "NONE": "",
        "PENDING": "Chờ duyệt",
        "APPROVED": "Đã duyệt",
        "REJECTED": "Từ chối",
    }.get(status, status)


MONTH_ABBR = (
    "JAN", "FEB", "MAR", "APR", "MAY", "JUN",
    "JUL", "AUG", "SEP", "OCT", "NOV", "DEC",
)

REPORT_AREA = "HEAD OFFICE"

# Ghi chú ký hiệu ngày công in ở đầu bảng chấm công tháng. Mỗi khối gồm tiêu
# đề, màu và các dòng (nhãn, ký hiệu); dòng có ký hiệu None là tiêu đề phụ.
NOTE_BLOCKS = (
    (
        "HƯỞNG LƯƠNG", NOTE_PAID_FILL, NOTE_PAID_ROW_FILL, (6, 13), (14, 19),
        (
            ("Ngày đi làm bình thường:", "HC"),
            ("Ngày nghỉ phép:", "Điền từ số 0.5 trở lên"),
            ("Kết hôn:", "KH"),
            ("Tang tế:", "TA"),
        ),
    ),
    (
        "CHẾ ĐỘ BHXH", NOTE_INSURANCE_FILL, NOTE_INSURANCE_ROW_FILL, (21, 28), (29, 30),
        (
            ("Bản thân ốm:", "OM"),
            ("Con ốm:", "CO"),
            ("Thai sản:", "TS"),
            ("Dưỡng sức sau khi sinh:", "DS"),
        ),
    ),
    (
        "CHẾ ĐỘ BHXH", NOTE_INSURANCE_FILL, NOTE_INSURANCE_ROW_FILL, (32, 36), (37, 37),
        (
            ("Khám thai", "KT"),
            ("Sẩy thai", "ST"),
            ("KHÔNG LƯƠNG", None),
            ("Nghỉ không lương:", "0"),
        ),
    ),
)


def _add_logo(sheet):
    """Đặt logo công ty ở góc trên bên trái, chiếm các dòng 1..LOGO_ROWS."""
    for row in range(1, LOGO_ROWS + 1):
        sheet.row_dimensions[row].height = 15
    if not LOGO_PATH.exists():
        return
    image = XLImage(str(LOGO_PATH))
    image.anchor = "A1"
    sheet.add_image(image)


def _styled_range(sheet, row, first_column, last_column, value, fill=None, font=None,
                  horizontal="left"):
    """Ghi một ô (gộp nếu trải nhiều cột) và kẻ viền cho mọi ô trong vùng gộp."""
    if last_column > first_column:
        sheet.merge_cells(
            start_row=row, start_column=first_column,
            end_row=row, end_column=last_column,
        )
    for column in range(first_column, last_column + 1):
        cell = sheet.cell(row=row, column=column)
        cell.border = GRID_BORDER
        if fill is not None:
            cell.fill = fill
    cell = sheet.cell(row=row, column=first_column, value=value)
    cell.font = font or Font(name="Times New Roman", size=10)
    cell.alignment = Alignment(horizontal=horizontal, vertical="center")
    return cell


def _add_timesheet_header(sheet, year, month):
    """Logo, thông tin bảng chấm công và khối ghi chú ký hiệu.

    Trả về số dòng khối tiêu đề chiếm, bảng dữ liệu bắt đầu ngay sau đó.
    """
    _add_logo(sheet)
    last_day = calendar.monthrange(year, month)[1]
    label_font = Font(name="Times New Roman", bold=True, size=11)

    info_rows = (
        (6, "BẢNG CHẤM CÔNG THÁNG", f"{month}/{year}"),
        (7, "KHU VỰC:", REPORT_AREA),
    )
    for row, label, value in info_rows:
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        label_cell = sheet.cell(row=row, column=1, value=label)
        label_cell.font = Font(
            name="Times New Roman", bold=True, size=11,
            color="FF0000" if row == 6 else "000000",
        )
        label_cell.alignment = Alignment(horizontal="left", vertical="center")
        value_cell = sheet.cell(row=row, column=3, value=value)
        value_cell.fill = SUBHEADER_FILL
        value_cell.font = label_font
        value_cell.alignment = Alignment(horizontal="center", vertical="center")

    # Cột C và E đủ rộng cho ngày; các cột ngày công chỉ rộng hơn 4 nên đặt
    # ngày vào đó sẽ hiện thành "#####".
    for row, label, value, second_label, second_value in (
        (8, "Từ ngày:", date(year, month, 1), "Đến ngày:", date(year, month, last_day)),
        (9, "Ngày xuất BC:", _report_date(), None, None),
    ):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        for column, content in (
            (1, label), (3, value), (4, second_label), (5, second_value)
        ):
            if content is None:
                continue
            cell = sheet.cell(row=row, column=column, value=content)
            cell.font = Font(name="Times New Roman", size=11, bold=column in {1, 4})
            cell.alignment = Alignment(
                horizontal="center" if column == 3 else "left", vertical="center"
            )
            if isinstance(content, date):
                cell.number_format = "dd/mm/yyyy"

    note_label = sheet.cell(row=6, column=4, value="GHI CHÚ:")
    note_label.font = Font(name="Times New Roman", bold=True, size=11, underline="single")
    note_label.alignment = Alignment(horizontal="right", vertical="center")

    title_row = 5
    for title, fill, row_fill, label_span, code_span, entries in NOTE_BLOCKS:
        _styled_range(
            sheet, title_row, label_span[0], code_span[1], title, fill,
            Font(name="Times New Roman", bold=True, size=10), "center",
        )
        for offset, (label, code) in enumerate(entries, start=1):
            row = title_row + offset
            if code is None:
                _styled_range(
                    sheet, row, label_span[0], code_span[1], label, fill,
                    Font(name="Times New Roman", bold=True, size=10), "center",
                )
                continue
            _styled_range(
                sheet, row, label_span[0], label_span[1], label, row_fill,
                Font(name="Times New Roman", italic=True, size=10),
            )
            _styled_range(
                sheet, row, code_span[0], code_span[1], code, None,
                Font(name="Times New Roman", bold=True, size=10),
            )

    for row in range(5, 10):
        sheet.row_dimensions[row].height = 16
    return 10


def _add_report_header(sheet, year, month):
    """Logo rồi tới tiêu đề "BẢNG CHẤM CÔNG AUG.26", kỳ báo cáo và ngày xuất.

    Trả về số dòng khối tiêu đề chiếm, bảng dữ liệu bắt đầu ngay sau đó.
    """
    _add_logo(sheet)
    last_day = calendar.monthrange(year, month)[1]
    title_row = LOGO_ROWS + 1

    sheet.merge_cells(
        start_row=title_row, start_column=1, end_row=title_row, end_column=4
    )
    title = sheet.cell(
        row=title_row,
        column=1,
        value=f"BẢNG CHẤM CÔNG {MONTH_ABBR[month - 1]}.{year % 100:02d}",
    )
    title.font = Font(name="Times New Roman", bold=True, size=14)
    title.alignment = Alignment(horizontal="left", vertical="center")

    # Nhãn nằm ở A:B đã gộp, ngày nằm ở cột C và E.
    info_rows = (
        (title_row + 1, "Từ ngày:", date(year, month, 1), "Đến ngày:", date(year, month, last_day)),
        (title_row + 2, "Ngày xuất BC:", _report_date(), None, None),
    )
    for row, label, value, second_label, second_value in info_rows:
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        for column, content in (
            (1, label), (3, value), (4, second_label), (5, second_value)
        ):
            if content is None:
                continue
            cell = sheet.cell(row=row, column=column, value=content)
            cell.font = Font(
                name="Times New Roman", size=11, bold=column in {1, 4}
            )
            cell.alignment = Alignment(horizontal="left", vertical="center")
            if isinstance(content, date):
                cell.number_format = "dd/mm/yyyy"

    sheet.row_dimensions[title_row].height = 22
    sheet.row_dimensions[title_row + 1].height = 18
    sheet.row_dimensions[title_row + 2].height = 18
    return title_row + 2


def _new_workbook(title):
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = title
    workbook.properties.creator = "Face Attendance"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    return workbook


def _save(workbook):
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def build_daily_workbook(rows, date_from=None, date_to=None):
    workbook = _new_workbook("Báo cáo chấm công")
    sheet = workbook.create_sheet("Cham cong")
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A5"

    title = "BÁO CÁO CHẤM CÔNG"
    if date_from and date_to and date_from == date_to:
        title += f" NGÀY {date_from.strftime('%d/%m/%Y')}"
    elif date_from:
        title += f" TỪ {date_from.strftime('%d/%m/%Y')}"
    elif date_to:
        title += f" ĐẾN {date_to.strftime('%d/%m/%Y')}"
    sheet.merge_cells("A1:M2")
    title_cell = sheet["A1"]
    title_cell.value = title
    title_cell.font = Font(name="Arial", size=16, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill("solid", fgColor="1F4E78")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 24
    sheet.row_dimensions[2].height = 12

    sheet["A3"] = "Thời điểm xuất:"
    sheet["B3"] = datetime.now()
    sheet["B3"].number_format = "dd/mm/yyyy hh:mm"
    sheet["A3"].font = Font(name="Arial", italic=True, color="666666")
    sheet["B3"].font = Font(name="Arial", italic=True, color="666666")

    headers = [
        "Mã nhân viên",
        "Họ tên",
        "Phòng ban",
        "Ca",
        "Ngày công",
        "Check-in",
        "Check-out",
        "Phút làm",
        "Đi trễ",
        "Về sớm",
        "Tăng ca",
        "Trạng thái",
        "Duyệt OT",
    ]
    for column, value in enumerate(headers, start=1):
        cell = sheet.cell(row=4, column=column, value=value)
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = GRID_BORDER
    sheet.row_dimensions[4].height = 30

    for row_index, (daily, employee, shift) in enumerate(rows, start=5):
        values = [
            _safe_text(employee.employee_code),
            _safe_text(employee.full_name),
            _department_name(employee),
            _safe_text(shift.code),
            daily.work_date,
            daily.first_check_in,
            daily.last_check_out,
            daily.worked_minutes,
            daily.late_minutes,
            daily.early_leave_minutes,
            AttendanceCalculator.format_duration(daily.overtime_minutes),
            _status_label(daily.status),
            _overtime_status_label(daily.overtime_status),
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_index, column=column, value=value)
            cell.font = Font(name="Arial", size=10)
            cell.border = GRID_BORDER
            cell.alignment = Alignment(
                horizontal="left" if column in {1, 2, 3, 4, 12} else "center",
                vertical="center",
            )
        sheet.cell(row=row_index, column=5).number_format = "dd/mm/yyyy"
        sheet.cell(row=row_index, column=6).number_format = "dd/mm/yyyy hh:mm"
        sheet.cell(row=row_index, column=7).number_format = "dd/mm/yyyy hh:mm"
        if row_index % 2 == 0:
            for column in range(1, 14):
                sheet.cell(row=row_index, column=column).fill = PatternFill(
                    "solid", fgColor="EAF2F8"
                )

    last_row = max(4, 4 + len(rows))
    sheet.auto_filter.ref = f"A4:M{last_row}"
    widths = [16, 28, 20, 12, 14, 20, 20, 12, 11, 11, 11, 20, 14]
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "1:4"
    return _save(workbook)


def build_monthly_workbook(
    employees,
    attendance_rows,
    year,
    month,
    half_day_ratio=0.75,
):
    workbook = _new_workbook(f"Báo cáo công tháng {month}/{year}")
    sheet = workbook.create_sheet(f"Thang {month:02d}-{year}")
    sheet.sheet_view.showGridLines = False

    # Bảng nằm dưới khối tiêu đề: hai dòng tiêu đề cột rồi mới tới nhân viên.
    top = _add_timesheet_header(sheet, year, month)
    head_row, day_row, first_data_row = top + 1, top + 2, top + 3
    sheet.freeze_panes = f"F{first_data_row}"

    last_day = calendar.monthrange(year, month)[1]
    employee_rows = list(employees)
    attendance = {
        (daily.employee_id, daily.work_date.day): (daily, shift)
        for daily, _employee, shift in attendance_rows
    }

    headers = ["STT", "MÃ NHÂN VIÊN", "TÊN NHÂN VIÊN", "PHÒNG BAN", "CA LÀM"]
    for column, value in enumerate(headers, start=1):
        sheet.merge_cells(
            start_row=head_row, start_column=column,
            end_row=day_row, end_column=column,
        )
        cell = sheet.cell(row=head_row, column=column, value=value)
        cell.fill = HEADER_FILL
        cell.font = Font(name="Times New Roman", bold=True, size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = GRID_BORDER

    sheet.merge_cells(
        start_row=head_row, start_column=6, end_row=head_row, end_column=36
    )
    month_cell = sheet.cell(row=head_row, column=6, value=f"THÁNG {month}/{year}")
    month_cell.fill = HEADER_FILL
    month_cell.font = Font(name="Times New Roman", bold=True, size=11)
    month_cell.alignment = Alignment(horizontal="center", vertical="center")
    month_cell.border = GRID_BORDER

    sheet.merge_cells(
        start_row=head_row, start_column=37, end_row=day_row, end_column=37
    )
    total_header = sheet.cell(row=head_row, column=37, value="TỔNG NGÀY\nCÔNG")
    total_header.fill = HEADER_FILL
    total_header.font = Font(name="Times New Roman", bold=True, size=10)
    total_header.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    total_header.border = GRID_BORDER

    for day in range(1, 32):
        column = day + 5
        cell = sheet.cell(row=day_row, column=column, value=day)
        valid_day = day <= last_day
        is_sunday = valid_day and date(year, month, day).weekday() == 6
        cell.fill = (
            WEEKEND_FILL if is_sunday else SUBHEADER_FILL if valid_day else INVALID_DAY_FILL
        )
        cell.font = Font(
            name="Times New Roman",
            bold=True,
            size=9,
            color="FFFFFF" if is_sunday else "000000",
        )
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = GRID_BORDER

    for row_index, (employee, shift) in enumerate(employee_rows, start=first_data_row):
        info_values = [
            row_index - first_data_row + 1,
            _safe_text(employee.employee_code),
            _safe_text(employee.full_name),
            _department_name(employee),
            _safe_text(shift.code if shift else ""),
        ]
        for column, value in enumerate(info_values, start=1):
            cell = sheet.cell(row=row_index, column=column, value=value)
            cell.font = Font(
                name="Times New Roman",
                size=10,
                bold=column == 1,
            )
            cell.alignment = Alignment(
                horizontal="center" if column in {1, 5} else "left",
                vertical="center",
            )
            cell.border = GRID_BORDER

        for day in range(1, 32):
            column = day + 5
            cell = sheet.cell(row=row_index, column=column)
            cell.border = GRID_BORDER
            cell.font = Font(name="Times New Roman", size=9)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if day > last_day:
                cell.fill = INVALID_DAY_FILL
                continue
            cell.fill = WORKDAY_FILL
            item = attendance.get((employee.id, day))
            if item is None:
                continue
            daily, daily_shift = item
            code, _value = _workday_code(daily, daily_shift, half_day_ratio)
            cell.value = code
            if code in {"HC", "0.5HC"}:
                cell.fill = PRESENT_FILL
            elif code in {"TCO", "TCI"}:
                cell.fill = WARNING_FILL
            elif code == "V":
                cell.fill = ABSENT_FILL
            elif code == "NP":
                cell.fill = LEAVE_FILL

        total_cell = sheet.cell(row=row_index, column=37)
        total_cell.value = (
            f'=COUNTIF(F{row_index}:AJ{row_index},"HC")+'
            f'0.5*COUNTIF(F{row_index}:AJ{row_index},"0.5HC")'
        )
        total_cell.number_format = "0.0#"
        total_cell.font = Font(name="Times New Roman", bold=True, size=10)
        total_cell.alignment = Alignment(horizontal="center", vertical="center")
        total_cell.border = GRID_BORDER
        sheet.row_dimensions[row_index].height = 22

    sheet.row_dimensions[head_row].height = 25
    sheet.row_dimensions[day_row].height = 22
    # Cột B rộng đủ cho nhãn "BẢNG CHẤM CÔNG THÁNG" gộp ở A:B; cột ngày rộng
    # đủ cho ký hiệu dài nhất "0.5HC".
    widths = {1: 6, 2: 24, 3: 28, 4: 18, 5: 13, 37: 13}
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = width
    for column in range(6, 37):
        sheet.column_dimensions[get_column_letter(column)].width = 5.5

    legend_row = max(5, len(employee_rows) + 5) + top
    sheet.merge_cells(start_row=legend_row, start_column=1, end_row=legend_row, end_column=12)
    legend = sheet.cell(
        row=legend_row,
        column=1,
        value=(
            "Quy ước: HC = đủ công; 0.5HC = nửa công; TCO = thiếu check-out; "
            "TCI = thiếu check-in; V = vắng; NP = nghỉ phép"
        ),
    )
    legend.font = Font(name="Times New Roman", italic=True, size=10)
    legend.alignment = Alignment(horizontal="left", vertical="center")

    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = f"1:{day_row}"
    sheet.print_area = f"A1:AK{max(legend_row, len(employee_rows) + day_row)}"
    _add_overtime_sheet(workbook, attendance_rows, year, month)
    return _save(workbook)


OT_ROWS_PER_EMPLOYEE = 3


def _merge_block(sheet, first_row, last_row, column):
    """Gộp một cột qua các dòng của một nhân viên và kẻ viền toàn khối."""
    sheet.merge_cells(
        start_row=first_row, start_column=column, end_row=last_row, end_column=column
    )
    for row in range(first_row, last_row + 1):
        sheet.cell(row=row, column=column).border = GRID_BORDER


def _add_overtime_sheet(workbook, attendance_rows, year, month):
    sheet = workbook.create_sheet("Công OT")
    sheet.sheet_view.showGridLines = False
    top = _add_report_header(sheet, year, month)

    head_row, day_row, first_data_row = top + 1, top + 2, top + 3
    sheet.freeze_panes = f"F{first_data_row}"

    # Only overtime an admin has signed off on is billable, so the sheet is the
    # approved set rather than everything the clock produced.
    overtime_rows = [
        (daily, employee, shift)
        for daily, employee, shift in attendance_rows
        if daily.overtime_minutes > 0 and daily.overtime_status == "APPROVED"
    ]
    overtime_dates = sorted({daily.work_date for daily, _employee, _shift in overtime_rows})
    employee_data = {}
    overtime_days = {}
    for daily, employee, _shift in overtime_rows:
        employee_data.setdefault(employee.id, employee)
        overtime_days[(employee.id, daily.work_date)] = daily
    employees = sorted(
        employee_data.values(),
        key=lambda employee: employee.employee_code,
    )

    day_count = max(1, len(overtime_dates))
    time_column = 5
    first_day_column = 6
    last_day_column = first_day_column + day_count - 1
    total_column = last_day_column + 1
    note_column = total_column + 1

    # Cột E chỉ chứa nhãn Giờ vào/Giờ ra/Giờ OT của từng nhân viên nên không
    # cần tiêu đề.
    static_headers = ["STT", "Mã NV", "TÊN NHÂN VIÊN", "PHÒNG BAN", None]
    for column, value in enumerate(static_headers, start=1):
        sheet.merge_cells(
            start_row=head_row,
            start_column=column,
            end_row=day_row,
            end_column=column,
        )
        cell = sheet.cell(row=head_row, column=column, value=value)
        cell.fill = OT_INFO_FILL
        cell.font = Font(name="Times New Roman", bold=True, size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = GRID_BORDER

    if last_day_column > first_day_column:
        sheet.merge_cells(
            start_row=head_row,
            start_column=first_day_column,
            end_row=head_row,
            end_column=last_day_column,
        )
    date_group = sheet.cell(row=head_row, column=first_day_column, value="Ngày (tiếng)")
    date_group.fill = OT_DATE_GROUP_FILL
    date_group.font = Font(name="Times New Roman", bold=True, size=11)
    date_group.alignment = Alignment(horizontal="center", vertical="center")
    date_group.border = GRID_BORDER

    for offset in range(day_count):
        column = first_day_column + offset
        value = overtime_dates[offset] if offset < len(overtime_dates) else "-"
        cell = sheet.cell(row=day_row, column=column, value=value)
        cell.fill = OT_DATE_FILL
        cell.font = Font(name="Times New Roman", bold=True, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = GRID_BORDER
        if isinstance(value, date):
            cell.number_format = "dd/mm"

    for column, value, fill in (
        (total_column, "Total", OT_TOTAL_FILL),
        (note_column, "Note", OT_NOTE_FILL),
    ):
        sheet.merge_cells(
            start_row=head_row,
            start_column=column,
            end_row=day_row,
            end_column=column,
        )
        cell = sheet.cell(row=head_row, column=column, value=value)
        cell.fill = fill
        cell.font = Font(name="Times New Roman", bold=True, size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = GRID_BORDER

    if employees:
        first_letter = get_column_letter(first_day_column)
        last_letter = get_column_letter(last_day_column)
        for index, employee in enumerate(employees):
            # Mỗi nhân viên chiếm 3 dòng: giờ vào, giờ ra và số giờ OT theo ngày.
            in_row = first_data_row + index * OT_ROWS_PER_EMPLOYEE
            out_row, hours_row = in_row + 1, in_row + 2
            info_values = [
                index + 1,
                _safe_text(employee.employee_code),
                _safe_text(employee.full_name),
                _department_name(employee),
            ]
            for column, value in enumerate(info_values, start=1):
                _merge_block(sheet, in_row, hours_row, column)
                cell = sheet.cell(row=in_row, column=column, value=value)
                cell.font = Font(name="Times New Roman", size=11)
                cell.alignment = Alignment(
                    horizontal="center" if column in {1, 4} else "left",
                    vertical="center",
                )

            for row, label in (
                (in_row, "Giờ vào"), (out_row, "Giờ ra"), (hours_row, "Giờ OT")
            ):
                cell = sheet.cell(row=row, column=time_column, value=label)
                cell.font = Font(name="Times New Roman", size=11)
                if row != hours_row:
                    cell.fill = OT_PUNCH_FILL
                cell.alignment = Alignment(horizontal="left", vertical="center")
                cell.border = GRID_BORDER
                sheet.row_dimensions[row].height = 18

            for offset, overtime_date in enumerate(overtime_dates):
                column = first_day_column + offset
                daily = overtime_days.get((employee.id, overtime_date))
                values = (
                    (in_row, _clock(daily.first_check_in) if daily else None, "h:mm"),
                    (out_row, _clock(daily.last_check_out) if daily else None, "h:mm"),
                    # Bảng OT chỉ tính theo giờ tròn nên lưu thẳng số giờ đã làm
                    # tròn, ô vừa hiện đúng con số vừa cộng được bằng SUM.
                    (
                        hours_row,
                        _overtime_hours(daily.overtime_minutes) if daily else None,
                        "0",
                    ),
                )
                for row, value, number_format in values:
                    cell = sheet.cell(row=row, column=column, value=value)
                    cell.number_format = number_format
                    cell.font = Font(name="Times New Roman", size=11)
                    if row != hours_row:
                        cell.fill = OT_PUNCH_FILL
                    # Canh giữa cho thẳng cột với tiêu đề ngày phía trên.
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.border = GRID_BORDER

            _merge_block(sheet, in_row, hours_row, total_column)
            total_cell = sheet.cell(row=in_row, column=total_column)
            total_cell.value = f"=SUM({first_letter}{hours_row}:{last_letter}{hours_row})"
            total_cell.number_format = "0"
            total_cell.font = Font(name="Times New Roman", bold=True, size=11)
            total_cell.alignment = Alignment(horizontal="center", vertical="center")

            _merge_block(sheet, in_row, hours_row, note_column)
            note_cell = sheet.cell(row=in_row, column=note_column)
            note_cell.font = Font(name="Times New Roman", size=11)
        last_row = day_row + len(employees) * OT_ROWS_PER_EMPLOYEE
    else:
        last_row = first_data_row
        sheet.merge_cells(
            start_row=first_data_row,
            start_column=1,
            end_row=first_data_row,
            end_column=note_column,
        )
        empty_cell = sheet.cell(
            row=first_data_row, column=1, value="Không có dữ liệu tăng ca trong tháng"
        )
        empty_cell.font = Font(name="Times New Roman", italic=True, size=11, color="666666")
        empty_cell.alignment = Alignment(horizontal="center", vertical="center")
        empty_cell.border = GRID_BORDER
        sheet.row_dimensions[first_data_row].height = 28

    sheet.row_dimensions[head_row].height = 30
    sheet.row_dimensions[day_row].height = 24
    widths = {
        1: 7, 2: 18, 3: 32, 4: 22, time_column: 10,
        total_column: 14, note_column: 18,
    }
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = width
    # Cột ngày chỉ cần đủ cho tiêu đề "dd/mm" và số giờ tăng ca bên dưới; để
    # rộng hơn thì bảng bị kéo dài vô ích khi tháng có nhiều ngày OT.
    for column in range(first_day_column, last_day_column + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 8

    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = f"1:{day_row}"
    sheet.print_area = f"A1:{get_column_letter(note_column)}{last_row}"


def build_monthly_report(db, year, month, employee_id=None, department_id=None):
    start_date = date(year, month, 1)
    end_date = date(year, month, calendar.monthrange(year, month)[1])
    employees = AttendanceRepository.list_employees_for_report(
        db,
        employee_id=employee_id,
        department_id=department_id,
        date_from=start_date,
        date_to=end_date,
    )
    rows = AttendanceRepository.list_daily(
        db,
        date_from=start_date,
        date_to=end_date,
        employee_id=employee_id,
        department_id=department_id,
    )
    return build_monthly_workbook(
        employees,
        rows,
        year,
        month,
        half_day_ratio=SettingsService.get_float(db, "HALF_DAY_RATIO"),
    )
