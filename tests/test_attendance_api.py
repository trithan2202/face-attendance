import unittest
from datetime import date, datetime, time
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import backend.app.models  # noqa: F401
from backend.app.core.config import settings
from backend.app.core.database import Base, get_db
from backend.app.core.security import hash_password
from backend.app.main import app
from backend.app.models import (
    AdminUser,
    AttendanceAdjustment,
    AttendanceDaily,
    AttendanceEvent,
    Employee,
    Shift,
)


class AttendanceApiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.original_secret = settings.JWT_SECRET
        settings.JWT_SECRET = "attendance-test-secret-at-least-32-bytes"
        cls.engine = create_engine(
            "sqlite://",
            connect_args={"check_same_thread": False},
            poolclass=StaticPool,
        )
        Base.metadata.create_all(cls.engine)
        cls.Session = sessionmaker(bind=cls.engine, expire_on_commit=False)
        with cls.Session() as db:
            admin = AdminUser(
                username="attendance-admin",
                password_hash=hash_password("correct-password"),
                is_active=True,
            )
            shift = Shift(
                code="HC-ATT",
                name="Hành chính",
                start_time=time(8),
                end_time=time(17),
                late_grace_minutes=10,
                early_leave_grace_minutes=10,
                cross_day=False,
            )
            db.add_all([admin, shift])
            db.flush()
            employee = Employee(
                employee_code="NV-ATT",
                full_name="Attendance Test",
                department="QA",
                shift_id=shift.id,
                status="ACTIVE",
            )
            db.add(employee)
            db.flush()
            daily = AttendanceDaily(
                employee_id=employee.id,
                shift_id=shift.id,
                work_date=date(2026, 8, 13),
                first_check_in=datetime(2026, 8, 13, 8, 5),
                last_check_out=None,
                status="MISSING_CHECK_OUT",
            )
            second_daily = AttendanceDaily(
                employee_id=employee.id,
                shift_id=shift.id,
                work_date=date(2026, 8, 14),
                first_check_in=datetime(2026, 8, 14, 8, 0),
                last_check_out=datetime(2026, 8, 14, 19, 0),
                worked_minutes=660,
                overtime_minutes=120,
                status="PRESENT",
            )
            regular_daily = AttendanceDaily(
                employee_id=employee.id,
                shift_id=shift.id,
                work_date=date(2026, 8, 15),
                first_check_in=datetime(2026, 8, 15, 8, 0),
                last_check_out=datetime(2026, 8, 15, 17, 0),
                worked_minutes=540,
                overtime_minutes=0,
                status="PRESENT",
            )
            db.add_all([daily, second_daily, regular_daily])
            db.flush()
            db.add(
                AttendanceEvent(
                    employee_id=employee.id,
                    event_type="CHECK_IN",
                    recognized_at=datetime(2026, 8, 13, 8, 5),
                    similarity=0.9,
                    detection_score=0.99,
                )
            )
            db.commit()
            cls.daily_id = daily.id

        def override_db():
            with cls.Session() as db:
                yield db

        app.dependency_overrides[get_db] = override_db
        cls.client = TestClient(app)
        login = cls.client.post(
            "/api/auth/login",
            json={
                "username": "attendance-admin",
                "password": "correct-password",
            },
        )
        assert login.status_code == 200

    @classmethod
    def tearDownClass(cls):
        cls.client.close()
        app.dependency_overrides.clear()
        cls.engine.dispose()
        settings.JWT_SECRET = cls.original_secret

    def test_01_list_detail_adjust_and_audit(self):
        listed = self.client.get(
            "/api/admin/attendance?date_from=2026-08-13&date_to=2026-08-13&department=QA"
        )
        self.assertEqual(listed.status_code, 200)
        self.assertEqual(len(listed.json()), 1)

        exported = self.client.get(
            "/api/admin/attendance/export.xlsx?date_from=2026-08-13&date_to=2026-08-13"
        )
        self.assertEqual(exported.status_code, 200)
        self.assertIn(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            exported.headers["content-type"],
        )
        daily_workbook = load_workbook(BytesIO(exported.content))
        daily_sheet = daily_workbook["Cham cong"]
        self.assertEqual(daily_sheet["A5"].value, "NV-ATT")
        self.assertEqual(daily_sheet["L5"].value, "Thiếu check-out")

        detail = self.client.get(f"/api/admin/attendance/{self.daily_id}")
        self.assertEqual(detail.status_code, 200)
        self.assertEqual(len(detail.json()["events"]), 1)

        no_change = self.client.patch(
            f"/api/admin/attendance/{self.daily_id}",
            json={"reason": "Không có thay đổi"},
        )
        self.assertEqual(no_change.status_code, 422)

        adjusted = self.client.patch(
            f"/api/admin/attendance/{self.daily_id}",
            json={
                "last_check_out": "2026-08-13T18:00:00",
                "reason": "Bổ sung giờ checkout theo xác nhận quản lý",
            },
        )
        self.assertEqual(adjusted.status_code, 200)
        result = adjusted.json()
        self.assertEqual(result["status"], "PRESENT")
        self.assertEqual(result["worked_minutes"], 595)
        self.assertEqual(result["overtime_minutes"], 60)

        audit = self.client.get(
            f"/api/admin/attendance/{self.daily_id}/adjustments"
        )
        self.assertEqual(audit.status_code, 200)
        self.assertEqual(len(audit.json()), 1)
        self.assertEqual(audit.json()[0]["admin_username"], "attendance-admin")
        self.assertIsNone(audit.json()[0]["before_data"]["last_check_out"])
        self.assertEqual(
            audit.json()[0]["after_data"]["last_check_out"],
            "2026-08-13T18:00:00",
        )

        monthly = self.client.get(
            "/api/admin/attendance/export-monthly.xlsx?year=2026&month=8"
        )
        self.assertEqual(monthly.status_code, 200)
        monthly_workbook = load_workbook(BytesIO(monthly.content), data_only=False)
        self.assertEqual(
            monthly_workbook.sheetnames,
            ["Thang 08-2026", "Công OT"],
        )
        monthly_sheet = monthly_workbook["Thang 08-2026"]
        self.assertIn("F1:AJ1", {str(item) for item in monthly_sheet.merged_cells.ranges})
        self.assertEqual(monthly_sheet["B3"].value, "NV-ATT")
        self.assertEqual(monthly_sheet["R3"].value, "HC")
        self.assertEqual(monthly_sheet["R3"].fill.fgColor.rgb, "0057D16C")
        self.assertEqual(monthly_sheet["U2"].fill.fgColor.rgb, "00FF4D4D")
        self.assertEqual(
            monthly_sheet["AK3"].value,
            '=COUNTIF(F3:AJ3,"HC")+0.5*COUNTIF(F3:AJ3,"0.5HC")',
        )
        overtime_sheet = monthly_workbook["Công OT"]
        self.assertIn("E1:F1", {str(item) for item in overtime_sheet.merged_cells.ranges})
        self.assertEqual(overtime_sheet["A1"].value, "STT")
        self.assertEqual(overtime_sheet["D1"].value, "PHÒNG BAN")
        self.assertEqual(overtime_sheet["E1"].value, "Ngày (tiếng)")
        self.assertEqual(overtime_sheet["E2"].value.date(), date(2026, 8, 13))
        self.assertEqual(overtime_sheet["F2"].value.date(), date(2026, 8, 14))
        self.assertEqual(overtime_sheet["B3"].value, "NV-ATT")
        self.assertEqual(overtime_sheet["D3"].value, "QA")
        self.assertEqual(overtime_sheet["E3"].value, 1)
        self.assertEqual(overtime_sheet["F3"].value, 2)
        self.assertEqual(overtime_sheet["G3"].value, "=SUM(E3:F3)")
        self.assertEqual(overtime_sheet["H1"].value, "Note")
        self.assertEqual(overtime_sheet.max_column, 8)
        self.assertEqual(overtime_sheet["A1"].fill.fgColor.rgb, "005B9BD5")
        self.assertEqual(overtime_sheet["E1"].fill.fgColor.rgb, "00FFC000")
        self.assertEqual(overtime_sheet["E2"].fill.fgColor.rgb, "00FF0000")
        self.assertEqual(overtime_sheet["G1"].fill.fgColor.rgb, "0082DC8C")

    def test_02_invalid_time_rolls_back_without_audit(self):
        invalid = self.client.patch(
            f"/api/admin/attendance/{self.daily_id}",
            json={
                "last_check_out": "2026-08-13T07:00:00",
                "reason": "Dữ liệu sai để kiểm thử rollback",
            },
        )
        self.assertEqual(invalid.status_code, 422)
        with self.Session() as db:
            count = db.scalar(select(func.count(AttendanceAdjustment.id)))
            self.assertEqual(count, 1)


if __name__ == "__main__":
    unittest.main()
